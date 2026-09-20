"""Deterministic build of the Trust by Construction V19 manuscripts.

Usage: python build_release.py [output-directory]

One markdown source per document produces both artifacts:
  * DOCX via python-docx, saved with fixed zip entry metadata so the bytes are
    reproducible on an equivalent toolchain;
  * a paginated print HTML rendered to PDF by headless Chrome, so the PDF is a
    rendering of the same source rather than a separately edited document.

Figures are built by make_figures.py and are original drawings of the
architecture or of recorded mechanism traces. No figure reports a field
measurement.
"""
from pathlib import Path
import datetime
import hashlib
import html as htmllib
import json
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from content import LISTINGS, TABLES

BASE = Path(__file__).resolve().parent
OUT = (Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else BASE / "build")
OUT.mkdir(parents=True, exist_ok=True)
REFS = json.loads((BASE / "references.json").read_text())
ABSTRACT_REFS = [1, 4, 6, 9, 12, 14, 17, 31, 37]
BUILD_STAMP = datetime.datetime(2026, 9, 20, 12, 0, 0)
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
VENUE_LINE = ("UNU Macau AI Conference 2026 \u2014 \u201cAI \u00d7 Education: AI for Learning, "
              "Learning for AI\u201d  |  Artifact: github.com/genaiworks/fssai-ra")
ALT = {
    "fig1-enforcement-boundary.png":
        "Three columns left to right: a lower-trust intelligence layer, a trusted control "
        "plane, and institutional systems. A typed proposal crosses into the control plane and "
        "a validated effect crosses out of it; a dashed path directly from the intelligence "
        "layer to institutional systems is marked impossible.",
    "fig2-proof-before-power.png":
        "Five gates in sequence - declared ceiling, task contract, lease and attenuation, "
        "mediated read and typed effect, sealed release - above bars showing the authority "
        "remaining after each gate, each shorter than the one before.",
    "fig3-pattern-map.png":
        "Five lifecycle stages - admit, read, delegate, act, release - each listing the "
        "patterns enforced there, beside the five-step contract applied at every stage.",
    "fig4-agent-graph.png":
        "A coordinator and three agents in a row; each agent's own scope is allowed, while the "
        "composed path runs from a protected source to a public sink and is rejected before the "
        "last agent acts.",
    "fig5-authority-states.png":
        "Five authority states from normal to quarantined, with automatic triggers that move a "
        "workload rightward and a restoration arrow that requires named human authority.",
    "fig6-delivery-revocation.png":
        "A ten-byte artifact with three bytes delivered, and five independent interventions, "
        "each denying both the next chunk and bulk delivery while the cursor stays at three bytes.",
}


# --------------------------------------------------------------- source model
def parse(kind):
    """Split a manuscript into (block, marker) pairs with captions attached."""
    text = (BASE / (kind + ".md")).read_text()
    text = re.sub(r"^(#{2,3} [^\n]+)\n", r"\1\n\n", text, flags=re.M)
    blocks = text.strip().split("\n\n")
    items, skip = [], False
    for idx, block in enumerate(blocks):
        if skip:
            skip = False
            continue
        caption = blocks[idx + 1] if idx + 1 < len(blocks) else ""
        if block.startswith("@FIG:"):
            items.append(("figure", block[5:], caption))
            skip = True
        elif block.startswith("@TABLE:"):
            items.append(("table", block[7:], caption))
            skip = True
        elif block.startswith("@LISTING:"):
            items.append(("listing", block[9:], caption))
            skip = True
        elif block.startswith("# "):
            lines = block.splitlines()
            subtitle = lines[1] if len(lines) > 1 else ""
            items.append(("title", lines[0][2:], "" if subtitle == "@AUTHOR" else subtitle))
            if subtitle == "@AUTHOR":
                items.append(("author", "", ""))
        elif block == "@AUTHOR":
            items.append(("author", "", ""))
        elif block == "@REFERENCES":
            items.append(("references", "", ""))
        elif block.startswith("### "):
            items.append(("h2", block[4:], ""))
        elif block.startswith("## "):
            items.append(("h1", block[3:], ""))
        elif block.startswith("Scale intelligence."):
            items.append(("cadence", block, ""))
        elif block.startswith("Property 1"):
            items.append(("property", block, ""))
        elif block.startswith("O(s,m,r) ⊆") or block.startswith("L(P,f) ="):
            expr = ("O(s, m, r) ⊆ A(s)" if block.startswith("O") else
                    "L(P, f) = log₂ |{T(P, f, z) : z ∈ Z}|")
            items.append(("equation", expr, "1" if block.startswith("O") else "2"))
        else:
            items.append(("para", block, ""))
    return items


def selected_refs(full):
    return REFS if full else [REFS[i - 1] for i in ABSTRACT_REFS]


def reference_text(ref, num):
    title = ref["title"].rstrip(".")
    dot = "" if title.endswith("?") else "."
    return f"[{num}] {ref['authors']} ({ref['year']}). {title}{dot} {ref.get('venue', '')}."


def check_citations(items, full):
    body = " ".join(b for kind, b, cap in items if kind != "references"
                    for b in (b, cap))
    cited = set()
    for group in re.findall(r"\[([0-9,– -]+)\]", body):
        for piece in group.split(","):
            bounds = re.split("[–-]", piece.strip())
            cited.update(range(int(bounds[0]), int(bounds[-1]) + 1))
    expected = set(range(1, (len(REFS) if full else len(ABSTRACT_REFS)) + 1))
    assert cited == expected, sorted(cited ^ expected)


# --------------------------------------------------------------------- DOCX
def xml(tag, **attrs):
    element = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    return element


def setfont(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    if rpr.find(qn("w:rFonts")) is None:
        rpr.append(OxmlElement("w:rFonts"))
    fonts = rpr.find(qn("w:rFonts"))
    for attr in list(fonts.attrib):
        if attr.endswith("Theme"):
            del fonts.attrib[attr]
    for slot in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn("w:" + slot), name)
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)


def hyperlink(paragraph, url):
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    props.append(xml("color", val="000000"))
    run.append(props)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


def docx_table(doc, key):
    header, rows, widths = TABLES[key]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(xml(edge, val="single", sz=4, color="D9D9D9"))
    table._tbl.tblPr.append(borders)
    for idx, values in enumerate([header] + rows):
        row = table.rows[0] if idx == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(xml("cantSplit"))
        if idx == 0:
            row._tr.get_or_add_trPr().append(xml("tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margins.append(xml(side, w=85, type="dxa"))
            props.append(margins)
            props.append(xml("shd", fill="E9EEF2" if idx == 0 else "FFFFFF"))
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.04
            para.paragraph_format.keep_with_next = idx == 0
            run = para.add_run(value)
            run.font.size = Pt(10)
            run.bold = idx == 0


def docx_listing(doc, key):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.8)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right"]:
        borders.append(xml(edge, val="single", sz=4, color="D9D9D9"))
    table._tbl.tblPr.append(borders)
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.8)
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ["top", "left", "bottom", "right"]:
        margins.append(xml(side, w=110, type="dxa"))
    props.append(margins)
    props.append(xml("shd", fill="F7F8F9"))
    for index, line in enumerate(LISTINGS[key]):
        para = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.12
        run = para.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
        for slot in ("ascii", "hAnsi"):
            fonts.set(qn("w:" + slot), "Courier New")
        rpr.append(fonts)


def build_docx(kind, items, full, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.85)
    section.footer_distance = Inches(.32)
    setfont(doc.styles["Normal"], "Times New Roman", 11)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.08
    normal.space_after = Pt(6)
    normal.widow_control = True
    for name, size in [("Title", 17), ("Heading 1", 12), ("Heading 2", 11)]:
        setfont(doc.styles[name], "Times New Roman", size, True)
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(0 if name == "Title" else 11)
        fmt.space_after = Pt(5)
        fmt.keep_with_next = True
    setfont(doc.styles["Caption"], "Times New Roman", 9.5)
    doc.styles["Caption"].font.italic = False
    doc.styles.add_style("Bibliography", 1)
    setfont(doc.styles["Bibliography"], "Times New Roman", 9.5)
    doc.styles["Bibliography"].paragraph_format.space_after = Pt(2)
    doc.styles["Bibliography"].paragraph_format.line_spacing = 1.0
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    for kindtag, body, extra in items:
        if kindtag == "title":
            doc.add_paragraph(body, "Title")
            if extra and extra != "@AUTHOR":
                para = doc.add_paragraph(extra)
                para.runs[0].italic = True
                para.paragraph_format.keep_with_next = True
        elif kindtag == "author":
            para = doc.add_paragraph("Rachana (Gen) Srivastava")
            para.runs[0].bold = True
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.keep_with_next = True
            para = doc.add_paragraph("Enterprise Architect | AI Systems Researcher")
            para.runs[0].font.size = Pt(9.5)
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.keep_with_next = True
            para = doc.add_paragraph(VENUE_LINE + ("" if full else "  |  Extended abstract"))
            para.runs[0].font.size = Pt(9.5)
            para.paragraph_format.space_after = Pt(10)
            para.paragraph_format.keep_with_next = True
        elif kindtag == "h1":
            para = doc.add_paragraph(body, "Heading 1")
            if body == "References":
                para.paragraph_format.page_break_before = full
        elif kindtag == "h2":
            doc.add_paragraph(body, "Heading 2")
        elif kindtag == "cadence":
            para = doc.add_paragraph(body)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].bold = True
            para.runs[0].italic = True
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(8)
        elif kindtag == "property":
            para = doc.add_paragraph(body)
            para.runs[0].bold = True
            para.paragraph_format.keep_with_next = True
        elif kindtag == "equation":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(8)
            math = OxmlElement("m:oMath")
            run = OxmlElement("m:r")
            text = OxmlElement("m:t")
            text.text = body
            run.append(text)
            math.append(run)
            para._p.append(math)
            para.add_run("   (" + extra + ")")
        elif kindtag == "figure":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(BASE / body), width=Inches(6.2))
            doc.inline_shapes[-1]._inline.docPr.set("descr", ALT[body])
            doc.add_paragraph(extra, "Caption")
        elif kindtag == "table":
            para = doc.add_paragraph(extra, "Caption")
            para.paragraph_format.keep_with_next = True
            docx_table(doc, body)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kindtag == "listing":
            docx_listing(doc, body)
            para = doc.add_paragraph(extra, "Caption")
            para.paragraph_format.space_before = Pt(4)
        elif kindtag == "references":
            for number, ref in enumerate(selected_refs(full), 1):
                para = doc.add_paragraph(style="Bibliography")
                para.paragraph_format.left_indent = Inches(.28)
                para.paragraph_format.first_line_indent = Inches(-.28)
                para.add_run(reference_text(ref, number))
                if ref.get("url"):
                    para.add_run(" ")
                    hyperlink(para, ref["url"])
        else:
            para = doc.add_paragraph(body)
            if body.startswith("Keywords:"):
                para.runs[0].font.size = Pt(9.5)
            if body.startswith(("Let s denote", "Let Z be", "Property 1")):
                para.paragraph_format.keep_with_next = True

    doc.core_properties.author = "Rachana (Gen) Srivastava"
    doc.core_properties.title = ("Trust by Construction: Secure-by-Design Patterns for "
                                 "Agentic AI in Learning Institutions")
    doc.core_properties.subject = "V-19 | " + ("Full technical paper" if full
                                               else "Extended abstract")
    doc.core_properties.version = "19"
    doc.core_properties.created = doc.core_properties.modified = BUILD_STAMP
    temp = path.with_suffix(".tmp.docx")
    doc.save(temp)
    with zipfile.ZipFile(temp) as zin, \
            zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name in sorted(zin.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 20, 12, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, zin.read(name))
    temp.unlink()
    words = " ".join(node.text or "" for node in doc._element.iter(qn("w:t")))
    return words


# --------------------------------------------------------------------- HTML
CSS = """
@page { size: 8.5in 11in; margin: 0.72in 0.85in; }
html, body { margin: 0; padding: 0; }
body { font-family: "Times New Roman", Times, serif; font-size: 11pt; color: #000;
       line-height: 1.24; width: 6.8in; margin: 0 auto;
       -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { height: 9.56in; position: relative; overflow: hidden; page-break-after: always; }
.page:last-child { page-break-after: auto; }
.pbody { height: 9.28in; overflow: hidden; }
.pfoot { position: absolute; bottom: 0; left: 0; right: 0; height: 0.2in;
         text-align: center; font-size: 9.5pt; color: #000; }
h1.doctitle { font-size: 17pt; margin: 0 0 4pt 0; line-height: 1.16; }
p.subtitle { font-style: italic; margin: 0 0 6pt 0; }
p.byline { font-weight: bold; margin: 0 0 1pt 0; }
p.role { font-size: 9.5pt; margin: 0 0 1pt 0; }
p.venue { font-size: 9.5pt; margin: 0 0 9pt 0; }
h2 { font-size: 12pt; margin: 11pt 0 5pt 0; }
h3 { font-size: 11pt; margin: 10pt 0 4pt 0; }
p { margin: 0 0 6pt 0; text-align: justify; }
p.keywords { font-size: 9.5pt; }
p.eq { text-align: center; margin: 4pt 0 8pt 0; }
p.prop { font-weight: bold; margin: 6pt 0; }
p.cadence { text-align: center; font-weight: bold; font-style: italic; margin: 4pt 0 8pt 0; }
figure { margin: 4pt 0 6pt 0; }
figure img { width: 91.2%; display: block; margin: 0 auto; }
p.caption, figcaption { font-size: 9.5pt; margin: 3pt 0 8pt 0; text-align: left; }
table { border-collapse: collapse; width: 100%; margin: 0 0 8pt 0; }
th, td { border: 0.5pt solid #d9d9d9; padding: 3pt 4pt; font-size: 10pt;
         vertical-align: middle; text-align: left; }
th { background: #e9eef2; font-weight: bold; }
pre.listing { border: 0.5pt solid #d9d9d9; background: #f7f8f9; margin: 0 0 3pt 0;
              padding: 6pt 8pt; font-family: "Courier New", monospace; font-size: 8.5pt;
              line-height: 1.32; white-space: pre; }
p.ref { font-size: 9.5pt; margin: 0 0 1.5pt 0; padding-left: 0.28in;
        text-indent: -0.28in; line-height: 1.12; text-align: left; }
a { color: #000; text-decoration: none; }
.keep { break-inside: avoid; }
"""

PAGINATE = """
function paginate() {
  var flow = document.getElementById('flow');
  var blocks = Array.prototype.slice.call(flow.children);
  var root = document.getElementById('pages');
  var limit = null, page = null, body = null, n = 0;
  function newPage() {
    n += 1;
    page = document.createElement('div'); page.className = 'page';
    body = document.createElement('div'); body.className = 'pbody';
    var foot = document.createElement('div'); foot.className = 'pfoot'; foot.textContent = n;
    page.appendChild(body); page.appendChild(foot); root.appendChild(page);
    if (limit === null) { limit = body.clientHeight; }
  }
  newPage();
  for (var i = 0; i < blocks.length; i++) {
    var block = blocks[i];
    body.appendChild(block);
    if (body.scrollHeight > limit) {
      var previous = block.previousElementSibling;
      var carry = null;
      if (previous && previous.dataset && previous.dataset.heading === '1'
          && previous.previousElementSibling) { carry = previous; }
      newPage();
      if (carry) { body.appendChild(carry); }
      body.appendChild(block);
    }
  }
  flow.remove();
  document.body.dataset.pages = n;
}
window.addEventListener('load', paginate);
"""


def esc(text):
    return htmllib.escape(text, quote=False)


def html_document(kind, items, full):
    out = []
    for kindtag, body, extra in items:
        if kindtag == "title":
            out.append(f'<h1 class="doctitle">{esc(body)}</h1>')
            if extra and extra != "@AUTHOR":
                out.append(f'<p class="subtitle">{esc(extra)}</p>')
        elif kindtag == "author":
            out.append('<p class="byline">Rachana (Gen) Srivastava</p>')
            out.append('<p class="role">Enterprise Architect | AI Systems Researcher</p>')
            out.append('<p class="venue">' + esc(VENUE_LINE)
                       + ("" if full else "  |  Extended abstract") + "</p>")
        elif kindtag == "h1":
            out.append(f'<h2 data-heading="1">{esc(body)}</h2>')
        elif kindtag == "h2":
            out.append(f'<h3 data-heading="1">{esc(body)}</h3>')
        elif kindtag == "cadence":
            out.append(f'<p class="cadence">{esc(body)}</p>')
        elif kindtag == "property":
            out.append(f'<p class="prop" data-heading="1">{esc(body)}</p>')
        elif kindtag == "equation":
            out.append(f'<p class="eq">{esc(body)} &nbsp;&nbsp;&nbsp;({extra})</p>')
        elif kindtag == "figure":
            out.append(f'<figure class="keep"><img src="{body}" alt="{esc(ALT[body])}">'
                       f"<figcaption>{esc(extra)}</figcaption></figure>")
        elif kindtag == "table":
            header, rows, _ = TABLES[body]
            cells = "".join(f"<th>{esc(c)}</th>" for c in header)
            table = [f"<table><thead><tr>{cells}</tr></thead><tbody>"]
            for row in rows:
                table.append("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in row) + "</tr>")
            table.append("</tbody></table>")
            out.append(f'<div class="keep"><p class="caption">{esc(extra)}</p>'
                       + "".join(table) + "</div>")
        elif kindtag == "listing":
            lines = "\n".join(esc(line) for line in LISTINGS[body])
            out.append(f'<div class="keep"><pre class="listing">{lines}</pre>'
                       f'<p class="caption">{esc(extra)}</p></div>')
        elif kindtag == "references":
            for number, ref in enumerate(selected_refs(full), 1):
                line = esc(reference_text(ref, number))
                if ref.get("url"):
                    line += f' <a href="{ref["url"]}">{esc(ref["url"])}</a>'
                out.append(f'<p class="ref">{line}</p>')
        else:
            css = ' class="keywords"' if body.startswith("Keywords:") else ""
            out.append(f"<p{css}>{esc(body)}</p>")
    title = "Trust by Construction V19" + ("" if full else " Extended Abstract")
    return ("<!doctype html><html><head><meta charset=\"utf-8\">"
            f"<title>{title}</title><style>{CSS}</style>"
            f"<script>{PAGINATE}</script></head><body>"
            f'<div id="pages"></div><div id="flow">' + "".join(out)
            + "</div></body></html>")


def render_pdf(html_path, pdf_path):
    if not Path(CHROME).exists():
        print("  ! Chrome not found; skipping PDF for", pdf_path.name)
        return False
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-sandbox",
                    "--no-pdf-header-footer", "--virtual-time-budget=20000",
                    "--run-all-compositor-stages-before-draw",
                    f"--print-to-pdf={pdf_path}", html_path.as_uri()],
                   check=True, capture_output=True)
    return pdf_path.exists()


# --------------------------------------------------------------------- build
def build(kind):
    full = kind == "full-paper"
    items = parse(kind)
    check_citations(items, full)
    stem = ("Trust_by_Construction_V19" if full
            else "Trust_by_Construction_V19_Extended_Abstract")
    docx_path = OUT / (stem + ".docx")
    words = build_docx(kind, items, full, docx_path)
    html_path = OUT / (stem + ".html")
    html_path.write_text(html_document(kind, items, full))
    for figure in {body for tag, body, _ in items if tag == "figure"}:
        shutil.copyfile(BASE / figure, OUT / figure)
    pdf_path = OUT / (stem + ".pdf")
    rendered = render_pdf(html_path, pdf_path)
    body_words = words.split("References")[0].split()
    info = {
        "document": stem,
        "docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
        "pdf_rendered": bool(rendered),
        "body_words_including_title_headings_captions_tables": len(body_words),
        "total_words_including_references": len(words.split()),
        "references": len(selected_refs(full)),
        "figures": sum(1 for tag, _, _ in items if tag == "figure"),
        "tables": sum(1 for tag, _, _ in items if tag == "table"),
        "occurrences_of_synthetic": words.lower().count("synthetic"),
    }
    print(f"  built {stem}: {info['total_words_including_references']} words, "
          f"{info['figures']} figures, {info['tables']} tables, "
          f"{info['references']} references")
    return info


if __name__ == "__main__":
    subprocess.run([sys.executable, str(BASE / "make_figures.py")], check=True)
    stats = [build("full-paper"), build("extended-abstract")]
    (BASE / "release-stats.json").write_text(json.dumps(stats, indent=2) + "\n")
    print(json.dumps(stats, indent=2))
