from pathlib import Path
import re, hashlib, json, zipfile, datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
import subprocess
W=Path(__file__).resolve().parent
REPO=W.parents[2]
O=Path(__import__('sys').argv[1]).resolve() if len(__import__('sys').argv)>1 else W.parent
O.mkdir(parents=True,exist_ok=True)
s=(W/'manuscript.md').read_text()
# An original vector diagram, also rasterised at 300 dpi for portable Word embedding.
fig=W/'swarm-architecture.pdf'; c=canvas.Canvas(str(fig), pagesize=(612,234),invariant=1)
navy=HexColor('#24465E'); gray=HexColor('#52606B'); blue=HexColor('#EFF5F9'); orange=HexColor('#A65328')
def txt(x,y,text,size=10,bold=False,color=navy):
 c.setFillColor(color);c.setFont('Helvetica-Bold' if bold else 'Helvetica',size);c.drawCentredString(x,y,text)
def box(x,y,w,h,fill=blue):
 c.setStrokeColor(navy);c.setFillColor(fill);c.setLineWidth(.8);c.roundRect(x,y,w,h,5,fill=1)
def arrow(x1,y1,x2,y2):
 c.setStrokeColor(navy);c.setFillColor(navy);c.setLineWidth(1);c.line(x1,y1,x2,y2)
 import math
 a=math.atan2(y2-y1,x2-x1);p=c.beginPath();p.moveTo(x2,y2);p.lineTo(x2-5*math.cos(a-.45),y2-5*math.sin(a-.45));p.lineTo(x2-5*math.cos(a+.45),y2-5*math.sin(a+.45));p.close();c.drawPath(p,fill=1,stroke=0)
txt(306,218,'ONE TASK ENVELOPE ACROSS PARALLEL WORKERS',12,True)
box(196,162,223,40)
txt(307.5,187,'Current trusted state',10,True)
txt(307.5,172,'Rights  |  Shared budget  |  Revocation epoch',9)
arrow(307.5,162,307.5,149)
box(5,33,163,167,HexColor('#FBF4EE'))
txt(86.5,184,'UNTRUSTED AGENTS',10,True,color=orange)
box(17,136,139,30,HexColor('#FFFFFF'));txt(86.5,147,'Coordinator proposes work',9)
for y,label in [(103,'Worker A'),(70,'Worker B')]:
 box(37,y,99,24,HexColor('#FFFFFF'));txt(86.5,y+8,label,9)
txt(86.5,46,'No infrastructure credentials',8.5,color=gray)
box(196,33,223,116)
txt(307.5,131,'ENFORCEMENT GATE',11,True)
for y,line in [(111,'Admit graph and reserve budget'),(94,'Mediate context and artifact handoffs'),(77,'Check lineage and exact operation'),(60,'Revalidate epoch and recipient'),(43,'Commit or reconcile; record receipt')]:txt(307.5,y,line,9)
box(451,33,156,167,HexColor('#F4F6F6'))
txt(529,183,'PROTECTED SYSTEMS',10,True)
for y,line in [(151,'Records and memory'),(128,'Versioned corrections'),(105,'Approved recipient')]:txt(529,y,line,9.5)
txt(529,71,'Only mediated access',9,True)
txt(529,49,'No public release by default',8.5,color=gray)
arrow(168,91,196,91);arrow(419,91,451,91)
txt(306,11,'Unknown state, exhausted budget or stale epoch: refuse the protected operation.',9,True,color=orange)
c.save()
subprocess.run(['pdftoppm','-singlefile','-r','300','-png',str(fig),str(W/'swarm-architecture')],check=True)
# The composition diagram: the prohibited path the paper is about.
c=canvas.Canvas(str(W/'composition.pdf'),pagesize=(612,132),invariant=1)
txt(306,116,'INDIVIDUAL PERMISSION DOES NOT AUTHORISE THE WHOLE PATH',11,True)
for x,title,sub in [(5,'Reader','Protected transcript'),(213,'Summariser','Derived student data'),(421,'Publisher','Public destination')]:
 box(x,43,186,54,HexColor('#F4F6F6'))
 txt(x+93,77,title,11,True);txt(x+93,58,sub,10)
arrow(191,69,213,69);arrow(399,69,421,69)
c.setStrokeColor(orange);c.setLineWidth(2.5);c.line(405,58,417,80);c.line(405,80,417,58)
txt(306,22,'Source restrictions persist through summarisation. The public release is refused.',10,True,color=orange)
c.save()
subprocess.run(['pdftoppm','-singlefile','-r','300','-png',str(W/'composition.pdf'),str(W/'composition')],check=True)
canonical=REPO/'evaluation'/'results'/'v1.0.0-domain-pack-matrix.json'
assert (W/'domain-pack-matrix.json').read_bytes()==canonical.read_bytes(), 'evidence copy has drifted from '+str(canonical)

# The claim matrix. Every BUILT row names a test that must exist and pass.
CLAIMS=[('Delegation cannot enlarge authority','Attenuated lease; ancestor intersection','BUILT','test_lease_cannot_expand_any_authority_axis'),
 ('Parallel siblings cannot multiply a budget','One task envelope; atomic reservation','BUILT','test_aggregate_sibling_budget_does_not_multiply'),
 ('Separate processes cannot race a commit','Serialised authorisation and commit','BUILT','test_independent_processes_serialize_budget_and_release'),
 ('Composition cannot erase a restriction','Bound source labels; recipient checks','BUILT','test_potential_source_to_sink_path_denied_even_for_clean_message'),
 ('Memory cannot carry a restriction away','Reauthorised memory; restart persistence','BUILT','test_restart_preserves_memory_restrictions_budget_and_rejects_legacy_adapter'),
 ('Revocation invalidates in-flight authority','Task epoch revalidated at every use','BUILT','test_effect_approval_substitution_and_guardian_epoch_invalidation'),
 ('A release cannot resume after revocation','Per-chunk reauthorisation','BUILT','test_stream_reauthorizes_every_chunk_and_bulk_cannot_bypass'),
 ('Checking each step is not checking the sequence','Stateful harness against a reference model','BUILT','test_stateful_sequences_hold_and_find_counterexamples_when_a_control_is_removed'),
 ('The same bound holds across distributed workers','Shared epoch and reservations','PROPOSED','Evaluation plan, section 4'),
 ('A delivered external effect can be undone','Adapter reconciliation only','NOT CLAIMED','Stated limit')]
json.dump([{'claim':c,'mechanism':m,'status':st,'evidence':e} for c,m,st,e in CLAIMS],open(W/'claim-matrix.json','w'),indent=2)
def wrap(text,font,size,width):
 lines=['']
 for token in re.split(r'(?<=[_ ])',text):
  if lines[-1] and c.stringWidth(lines[-1]+token,font,size)>width:lines.append(token)
  else:lines[-1]+=token
 return [line.rstrip() for line in lines]
X=[8,225,0,455];WIDTH=[210,180,0,149];ROW=20.5;TOP=64+ROW*len(CLAIMS)
green=HexColor('#1F6F4A');amber=HexColor('#A65328')
c=canvas.Canvas(str(W/'claim-matrix.pdf'),pagesize=(612,TOP),invariant=1)
txt(306,TOP-17,'WHAT IS BUILT AND WHAT IS PROPOSED',12,True)
c.setFillColor(gray);c.setFont('Helvetica-Bold',7.2)
for x,head in [(X[0],'CLAIM'),(X[1],'MECHANISM THAT ENFORCES IT'),(X[3],'EVIDENCE IN THE REFERENCE KERNEL')]:c.drawString(x,TOP-34,head)
c.drawCentredString(415,TOP-34,'STATUS')
c.setStrokeColor(HexColor('#C8D2D8'));c.setLineWidth(.7);c.line(8,TOP-40,604,TOP-40)
for i,(claim,mech,status,ev) in enumerate(CLAIMS):
 top=TOP-46-i*ROW
 if i%2:c.setFillColor(HexColor('#F7F9FA'));c.rect(8,top-ROW+4,596,ROW,fill=1,stroke=0)
 c.setFillColor(navy);c.setFont('Helvetica-Bold',8.2);c.drawString(X[0],top-9,claim)
 c.setFillColor(gray);c.setFont('Helvetica',8);c.drawString(X[1],top-9,mech)
 c.setFillColor({'BUILT':green,'PROPOSED':amber}.get(status,gray));c.setFont('Helvetica-Bold',7.2)
 c.drawCentredString(415,top-9,status)
 c.setFillColor(gray);c.setFont('Helvetica-Oblique',6.8)
 for j,line in enumerate(wrap(ev,'Helvetica-Oblique',6.8,WIDTH[3])[:2]):c.drawString(X[3],top-6-j*7.6,line)
c.line(8,TOP-46-len(CLAIMS)*ROW+4,604,TOP-46-len(CLAIMS)*ROW+4)
c.setFillColor(gray);c.setFont('Helvetica',7.4)
c.drawCentredString(306,9,'Every BUILT row names an offline regression test. PROPOSED rows are obligations this abstract specifies and does not evaluate.')
subprocess.run(['pdftoppm','-singlefile','-r','300','-png',str(W/'claim-matrix.pdf'),str(W/'claim-matrix')],check=True)
c=canvas.Canvas(str(W/'patterns.pdf'),pagesize=(612,220),invariant=1)
txt(306,205,'TEN PATTERNS ON ONE TASK LIFECYCLE',12,True)
cols=[('ADMIT',[('P1','Declared ceiling'),('P2','Task contract'),('P3','Credentialless runtime')]),('READ',[('P5','Mediated context'),('P6','Reauthorised memory')]),('DELEGATE',[('P4','Attenuated delegation'),('P10','Declared graph')]),('ACT',[('P7','Typed effect')]),('RELEASE',[('P8','Sealed release'),('P9','Restricting monitor')])]
for i,(stage,rows) in enumerate(cols):
 x=5+i*122;box(x,61,114,129)
 c.setFillColor(HexColor('#2A78D6'));c.rect(x+3,166,108,21,fill=1,stroke=0);txt(x+57,172,stage,11,True,color=HexColor('#FFFFFF'))
 for j,(pid,label) in enumerate(rows):
  y=127-j*31;box(x+4,y,106,29,HexColor('#FFFFFF'));txt(x+57,y+17,pid,10,True);txt(x+57,y+5,label,9)
 if i<4:arrow(x+114,114,x+122,114)
box(5,5,602,46,HexColor('#F4F6F6'))
txt(306,36,'REPEAT AT EVERY BOUNDARY',10,True)
txt(306,21,'Authenticate  >  Load current state  >  Intersect rights  >  Validate  >  Commit safely',10)
txt(306,9,'Refuse on failed checks. Monitoring and graph checks are not limited to their displayed stage.',8.5)
c.save()
subprocess.run(['pdftoppm','-singlefile','-r','300','-png',str(W/'patterns.pdf'),str(W/'patterns')],check=True)

DISPLAY={'Introduction':'1. Authority before autonomy',
 'Development Section 1 Methodology Core Argument and Case Context':'2. Ten patterns, one boundary, and a swarm contract',
 'Development Section 2 Results Analysis and Impact':'3. Evidence, falsification and impact',
 'Conclusion':'4. Conclusion: scale intelligence, bound authority',
 'References':'References'}
ref=W/'page-system.docx'  # the V19 page system, vendored so this directory rebuilds standalone
assert hashlib.sha256(ref.read_bytes()).hexdigest()=='d535a59408f527f790d527d05c938b9bf147466d3f2802c2ef72051c37f4ec09'
doc=Document(ref)
for child in list(doc._element.body):
 if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
# Retain V19 page system and styles. All text/figure slots are revised.
for i,block in enumerate(s.strip().split('\n\n')):
 if block.startswith('# '):
  title,subtitle=block.split('\n',1);doc.add_paragraph(title[2:],'Title');p=doc.add_paragraph(subtitle);p.runs[0].italic=True;p.paragraph_format.keep_with_next=True
 elif block.startswith('Rachana'):
  for j,line in enumerate(block.splitlines()):
   p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=True
   p.runs[0].bold=j==0
   if j:p.runs[0].font.size=Pt(9.5)
 elif block.startswith('## '):
  p=doc.add_paragraph(DISPLAY.get(block[3:],block[3:]),'Heading 1')
 elif block.startswith('@FIG:'):
  figure=block.split(':',1)[1]
  p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
  r=p.add_run();r.add_picture(str(W/(figure+'.png')),width=Inches({'patterns':6.5,'swarm-architecture':6.4,'composition':6.4,'claim-matrix':6.6}[figure]))
  pr=r._r.xpath('.//wp:docPr')[0];pr.set('descr', {'patterns':'Ten patterns arranged across admit, read, delegate, act and release, with checks repeated at every boundary.', 'swarm-architecture':'Untrusted workers share one task envelope and access protected systems through an enforcement gate.', 'composition':'A reader, summariser and publisher form a protected-data path; the public release is refused.', 'claim-matrix':'A table of ten claims. Eight are built and name a regression test, one distributed bound is proposed and not evaluated, and undoing a delivered external effect is not claimed.'}[figure])
 elif re.match(r'^Figure \d+\.',block):
  p=doc.add_paragraph(block,'Caption');p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_with_next=False
 elif re.match(r'^\[\d+\]',block):
  p=doc.add_paragraph(block,'Bibliography');p.paragraph_format.keep_together=True
 else:
  p=doc.add_paragraph(block);p.paragraph_format.widow_control=True
# Consistent metadata and refreshable page fields.
doc.core_properties.title=s.splitlines()[0][2:];doc.core_properties.author='Rachana Srivastava';doc.core_properties.subject='UNU Macau AI Conference 2026 extended abstract';doc.core_properties.comments=''
STAMP=datetime.datetime(2026,9,20,12,0,0)
doc.core_properties.created=doc.core_properties.modified=STAMP
doc.core_properties.revision=1
out=O/'Trust_by_Construction_V24_Extended_Abstract.docx'
tmp=out.with_suffix('.tmp.docx');doc.save(tmp)
with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as zout:
 for name in sorted(zin.namelist()):
  info=zipfile.ZipInfo(name,date_time=(2026,9,20,12,0,0));info.compress_type=zipfile.ZIP_DEFLATED
  zout.writestr(info,zin.read(name))
tmp.unlink()
# Canonical editable text with image link, suitable for submission forms.
md=re.sub(r'@FIG:([a-z-]+)',r'![Figure](source/\1.png)',s)
(O/'Trust_by_Construction_V24_Extended_Abstract.md').write_text(md)
# The form takes pasted text and has no image field, so the paste drops figures and captions.
paste='\n'.join(l for l in s.splitlines() if not l.startswith('@FIG:') and not re.match(r'^Figure \d+\.',l))
paste=re.sub(r'\n{3,}','\n\n',paste)
(O/'Trust_by_Construction_V24_Form_Fields.md').write_text(paste)
fields=O/'form-fields';fields.mkdir(exist_ok=True)
for old in fields.glob('*.txt'):old.unlink()
order=['Introduction','Development Section 1 Methodology Core Argument and Case Context',
       'Development Section 2 Results Analysis and Impact','Conclusion','References']
blocks=dict();current=None
for line in paste.splitlines():
 if line.startswith('## '):current=line[3:].strip();blocks[current]=[]
 elif current:blocks[current].append(line)
for i,name in enumerate(order,1):
 body='\n'.join(blocks.get(name,[])).strip()
 (fields/f"{i}-{re.sub(r'[^a-z0-9]+','-',name.lower()).strip('-')[:40]}.txt").write_text(body+'\n')
# The PDF a reviewer opens is rendered from the same DOCX this script just wrote.
WPS='/Applications/wpsoffice.app/Contents/MacOS/wpscli'
pdf=O/(out.stem+'.pdf')
if Path(WPS).exists():
 pdf.unlink(missing_ok=True)
 rendered=subprocess.run([WPS,'word2pdf','--input',str(out),'--output',str(O)],capture_output=True)
 if rendered.returncode:print('PDF not rendered here (exit %d); DOCX is unaffected.'%rendered.returncode)
prose=re.sub(r'@FIG:\S+','',s)
body=prose.split('## References')[0]
counts={'body_words_excluding_figure_text':len(re.sub(r'^Figure \d+\..*$','',body,flags=re.M).split()),
        'reference_entries':len(re.findall(r'^\[\d+\] ',s,re.M)),
        'figures_in_manuscript':len(re.findall(r'@FIG:',s))}
cited={int(n) for n in re.findall(r'\[(\d+)\]',body)}
assert cited=={i+1 for i in range(counts['reference_entries'])}, ('citations do not resolve',cited)
# The form's own limits decide validity, imported from the checklist's checker so there is one rule.
import sys as _sys; _sys.path.insert(0,str(REPO))
from scripts.check_submission import validate
form=validate((O/'Trust_by_Construction_V24_Form_Fields.md').read_text())
assert form['valid'], json.dumps(form,indent=2)
(O/'release-verification.json').write_text(json.dumps({
 'manuscript':'Trust_by_Construction_V24_Extended_Abstract',
 'author':'Rachana Srivastava',
 'one_source':'source/manuscript.md renders both the document and the form fields',
 **counts,
 'form_fields_valid':form['valid'],
 'form_fields':{name:{'words':d['words'],'characters':d['characters'],
                      'character_headroom':d['headroom']} for name,d in form['sections'].items()},
 'pages':int(re.search(r'^Pages:\s+(\d+)',subprocess.run(['pdfinfo',str(pdf)],capture_output=True,text=True).stdout,re.M).group(1)) if pdf.exists() else None,
 'evidence_matches_repository_artifact':'evaluation/results/v1.0.0-domain-pack-matrix.json',
 'sha256':{'docx':hashlib.sha256(out.read_bytes()).hexdigest(),
           'pdf':hashlib.sha256(pdf.read_bytes()).hexdigest() if pdf.exists() else None},
},indent=2)+'\n')
counts['form_fields_valid']=form['valid']
print(json.dumps(counts,indent=2));print(out)
